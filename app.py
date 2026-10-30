from datetime import datetime
import glob
import io
import os
import re
import subprocess
import duckdb
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import plotly.express as px
import streamlit as st

st.set_page_config(
    page_title="QC Accuracy Dashboard",
    page_icon="🔥",
    layout="wide",
)

st.markdown(
    """
    <style>
    .main .block-container {
        padding-top: 0.5rem !important;
        padding-bottom: 0.5rem !important;
    }

    div[data-testid="stVerticalBlock"] > div:has(div.sticky-header-anchor) {
        position: sticky;
        top: 2.875rem;
        background-color: #0d1117;
        z-index: 999;
        padding-top: 0.25rem;
        padding-bottom: 0.25rem;
        border-bottom: 1px solid #21262d;
        margin-bottom: 0.25rem;
    }

    div[data-testid="stMetric"] {
        background-color: #161b22;
        border: 1px solid #30363d;
        border-radius: 6px;
        padding: 4px 10px !important;
    }

    [data-testid="stMetricLabel"] {
        font-size: 0.75rem !important;
        color: #8b949e !important;
        margin-bottom: -2px !important;
    }

    [data-testid="stMetricValue"] {
        font-size: 1.1rem !important;
        line-height: 1.2 !important;
    }

    [data-testid="stMetricDelta"] {
        font-size: 0.75rem !important;
        color: #8b949e !important;
    }

    [data-testid="stMetricDelta"] svg {
        display: none !important;
    }

    .js-plotly-plot .plotly .draglayer,
    .js-plotly-plot .plotly .nsewdrag,
    div[data-testid="stPlotlyChart"] .main-svg {
        cursor: pointer !important;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

SNAPSHOT_DIR = "snapshots"
LATEST_CACHE_FILE = "qc_local_cache.duckdb"

HEADER_MAP = {
    "Sales API V4 - Quantity": "Quantity",
    "Sales API V4 - Revenue": "Revenue",
    "Sales SKU - Quantity": "Quantity",
    "Sales SKU - Revenue": "Revenue",
    "Sales Seller - Quantity": "Quantity",
    "Sales Seller - Revenue": "Revenue",
    "Traffic SKU - Page View": "Page View",
    "Traffic SKU - Visitor": "Visitor",
    "Traffic Seller - Page View": "Page View",
}


def get_available_snapshots():
    file_list = []

    if os.path.exists(LATEST_CACHE_FILE):
        mtime = os.path.getmtime(LATEST_CACHE_FILE)
        file_list.append((mtime, LATEST_CACHE_FILE))

    if os.path.exists(SNAPSHOT_DIR):
        files = glob.glob(os.path.join(SNAPSHOT_DIR, "qc_cache_*.duckdb"))
        for fpath in files:
            mtime = os.path.getmtime(fpath)
            file_list.append((mtime, fpath))

    file_list.sort(key=lambda x: x[0], reverse=True)

    snapshots = {}
    for idx, (mtime, fpath) in enumerate(file_list):
        dt_str = datetime.fromtimestamp(mtime).strftime("%Y-%m-%d %H:%M:%S")
        icon = "⚡" if idx == 0 else "📁"
        label = f"{icon} {dt_str}"

        while label in snapshots:
            label += " "

        snapshots[label] = fpath

    return snapshots


@st.cache_data(ttl=5)
def load_data_from_snapshot(file_path):
    if not os.path.exists(file_path):
        return pd.DataFrame(), None, 0

    mtime = os.path.getmtime(file_path)
    snapshot_time = datetime.fromtimestamp(mtime).strftime("%Y-%m-%d %H:%M:%S")
    file_size_mb = os.path.getsize(file_path) / (1024 * 1024)

    con = duckdb.connect(file_path, read_only=True)
    df = con.execute("SELECT * FROM qc_accuracy_daily").df()
    con.close()

    if not df.empty and "company_seller_used_id" in df.columns:
        df = df[
            df["company_seller_used_id"].notna()
            & (df["company_seller_used_id"].astype(str).str.strip() != "")
            ].copy()

    if "seller_name" not in df.columns:
        df["seller_name"] = ""

    return df, snapshot_time, file_size_mb


def determine_source(metric_name):
    metric_lower = str(metric_name).lower()

    if "api" in metric_lower or "item" in metric_lower:
        return "item"
    elif "export sku" in metric_lower or "sku_sales" in metric_lower:
        if "traffic" in metric_lower or "pageview" in metric_lower or "visitor" in metric_lower:
            return "export_sku_traffic"
        return "export_sku_sales"
    elif "export seller" in metric_lower or "seller_sales" in metric_lower or "sales" in metric_lower:
        if "traffic" in metric_lower or "pageview" in metric_lower:
            return "export_seller_traffic"
        return "export_seller_sales"

    return "export_seller_sales"


def extract_version(metric_name):
    match = re.search(r"v(\d+)", str(metric_name), re.IGNORECASE)
    if match:
        return f"v{match.group(1).lower()}"
    return "v4"


def style_excel_sheet(ws):
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"

    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            header_name = str(ws.cell(row=1, column=cell.column).value).lower()

            if isinstance(cell.value, (int, float)):
                if "revenue" in header_name:
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                if "day" in header_name or "version" in header_name or "source" in header_name:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if cell.number_format and isinstance(cell.value, float):
                val_str = f"{cell.value:,.2f}"
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def generate_dual_sheet_excel(df_source, is_full_company=False, metric_name_for_tab="sales_traffic"):
    output = io.BytesIO()

    if is_full_company:
        tab1_name = "automated_input"
        tab2_name = "database_input"
    else:
        clean_metric_scope = re.sub(r"[\[\]:*?/\\]", "", metric_name_for_tab).strip().replace(" ", "_").lower()
        tab1_name = f"{clean_metric_scope}_auto_input"[:31]
        tab2_name = f"{clean_metric_scope}_db_input"[:31]

    df_prep = df_source.copy()
    df_prep["day"] = pd.to_datetime(df_prep["day"]).dt.strftime("%Y-%m-%d")

    if is_full_company:
        df_auto_pivot = df_prep.pivot_table(
            index=["company_seller_used_id", "seller_name", "day"],
            columns="metric_name",
            values="raw_val",
            aggfunc="first",
        ).reset_index()

        df_auto_pivot.columns.name = None
        df_tab1_final = df_auto_pivot.copy()

        metric_cols = [c for c in df_tab1_final.columns if c not in ["company_seller_used_id", "seller_name", "day"]]
        rename_dict_t1 = {col: HEADER_MAP.get(col, col) for col in metric_cols}
        df_tab1_final = df_tab1_final.rename(columns=rename_dict_t1)

        df_tab1_final.insert(3, "# days have data", 1)

        if "source" not in df_prep.columns or df_prep["source"].isna().all():
            df_prep["source"] = df_prep["metric_name"].apply(determine_source)

        df_prep["version"] = df_prep["metric_name"].apply(extract_version)

        df_db_pivot = df_prep.pivot_table(
            index=["company_seller_used_id", "seller_name", "day", "source"],
            columns="metric_name",
            values="collected_val",
            aggfunc="first",
        ).reset_index()

        df_db_pivot.columns.name = None
        df_tab2_final = df_db_pivot.copy()

        db_metric_cols = [c for c in df_tab2_final.columns if
                          c not in ["company_seller_used_id", "seller_name", "day", "source"]]
        rename_dict_t2 = {col: HEADER_MAP.get(col, col) for col in db_metric_cols}
        df_tab2_final = df_tab2_final.rename(columns=rename_dict_t2)

        df_tab2_final.insert(4, "# days have data", 1)

        version_map = df_prep.groupby(["company_seller_used_id", "seller_name", "day", "source"])[
            "version"].first().reset_index()
        df_tab2_final = df_tab2_final.merge(version_map, on=["company_seller_used_id", "seller_name", "day", "source"],
                                            how="left")

    else:
        df_tab1_final = df_prep[["company_seller_used_id", "seller_name", "day", "raw_val"]].copy()
        df_tab1_final.insert(3, "# days have data", 1)

        short_header = HEADER_MAP.get(metric_name_for_tab, "Quantity")
        df_tab1_final = df_tab1_final.rename(columns={"raw_val": short_header})

        df_tab2 = df_prep.copy()
        if "source" not in df_tab2.columns or df_tab2["source"].isna().all():
            df_tab2["source"] = df_tab2["metric_name"].apply(determine_source)

        df_tab2["version"] = df_tab2["metric_name"].apply(extract_version)

        df_tab2_final = df_tab2[
            ["company_seller_used_id", "seller_name", "day", "source", "collected_val", "version"]].copy()
        df_tab2_final.insert(4, "# days have data", 1)
        df_tab2_final = df_tab2_final.rename(columns={"collected_val": short_header})

        cols = [c for c in df_tab2_final.columns if c != "version"] + ["version"]
        df_tab2_final = df_tab2_final[cols]

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_tab1_final.to_excel(writer, sheet_name=tab1_name, index=False)
        df_tab2_final.to_excel(writer, sheet_name=tab2_name, index=False)

        ws1 = writer.sheets[tab1_name]
        ws2 = writer.sheets[tab2_name]

        style_excel_sheet(ws1)
        style_excel_sheet(ws2)

    return output.getvalue()


filter_container = st.sidebar.container()
st.sidebar.markdown("---")
snapshot_container = st.sidebar.container()

available_snapshots = get_available_snapshots()

if not available_snapshots:
    st.error("⚠️ No snapshot files found! Please run the ETL pipeline first.")
    st.stop()

with snapshot_container.expander("📷 Audit Sync Snapshot", expanded=False):
    selected_snapshot_label = st.selectbox(
        "Select Snapshot to Audit",
        options=list(available_snapshots.keys()),
        help="Select a historical snapshot record to audit past data.",
    )

    selected_snapshot_path = available_snapshots[selected_snapshot_label]
    df, snapshot_time, file_size_mb = load_data_from_snapshot(selected_snapshot_path)

    if st.button("🔄 Trigger New Sync (fetch_data.py)", use_container_width=True):
        with st.spinner("Executing fetch_data.py to fetch the latest data..."):
            try:
                result = subprocess.run(["python", "fetch_data.py"], capture_output=True, text=True)
                if result.returncode == 0:
                    st.success("✅ Data sync successful!")
                    st.cache_data.clear()
                    st.rerun()
                else:
                    st.error(f"❌ ETL Error: {result.stderr}")
            except Exception as e:
                st.error(f"❌ Execution Error: {e}")

    st.caption(f"📦 File: `{os.path.basename(selected_snapshot_path)}` ({file_size_mb:.2f} MB)")

with filter_container:
    st.header("🔍 Data Filters")

    if df.empty:
        st.warning("⚠️ Selected snapshot contains no valid data.")
        st.stop()

    companies = sorted(df["company_name"].dropna().unique())

    with st.form(key="data_filter_form"):
        selected_company = st.selectbox("Select Company", companies)

        df_company = df[df["company_name"] == selected_company]
        metrics = sorted(df_company["metric_name"].dropna().unique())
        selected_metric = st.selectbox("Select Metric", metrics)

        df_metric = df_company[df_company["metric_name"] == selected_metric].copy()

        all_sellers = sorted(df_metric["company_seller_used_id"].dropna().unique())
        select_all_sellers = st.checkbox("Select All Sellers", value=False)

        default_sellers = all_sellers if select_all_sellers else (
            all_sellers[:5] if len(all_sellers) >= 5 else all_sellers)

        seller_name_dict = dict(zip(df_metric["company_seller_used_id"],
                                    df_metric["seller_name"])) if "seller_name" in df_metric.columns else {}


        def format_seller_display(seller_id_str):
            name = seller_name_dict.get(seller_id_str, "")
            last_4 = seller_id_str[-4:] if len(seller_id_str) >= 4 else seller_id_str
            if name and pd.notna(name) and str(name).strip() != "":
                return f"{name} ...{last_4}"
            return f"...{last_4}"


        selected_sellers_input = st.multiselect(
            "Select Seller ID",
            options=all_sellers,
            default=default_sellers,
            format_func=format_seller_display,
        )

        apply_filters = st.form_submit_button("🚀 Apply Filters", use_container_width=True, type="primary")

        if select_all_sellers:
            selected_sellers = all_sellers
        else:
            selected_sellers = selected_sellers_input

    df_display = df_metric[df_metric["company_seller_used_id"].isin(selected_sellers)].copy()

    st.markdown("---")
    st.subheader("📥 Export Excel (GG Sheet Format)")

    export_scope = st.radio(
        "Export Scope:",
        options=["Filtered Data Only", "Full Company Data"],
        index=0,
        help="Choose whether to export only currently filtered sellers/metric or all company data.",
    )

    df_export_source = df_display if export_scope == "Filtered Data Only" else df_company

    if not df_export_source.empty:
        is_full = (export_scope != "Filtered Data Only")

        excel_bytes = generate_dual_sheet_excel(
            df_export_source,
            is_full_company=is_full,
            metric_name_for_tab=selected_metric if not is_full else "full_company"
        )

        clean_company = selected_company.replace(" ", "_").lower()
        clean_metric = selected_metric.replace(" ", "_").replace("-", "").lower()
        time_str = datetime.now().strftime("%Y%m%d_%H%M%S")

        scope_prefix = "FULL" if is_full else "filter"
        export_filename = f"qc_audit_{scope_prefix}_{clean_company}_{time_str}.xlsx"

        st.download_button(
            label="📊 Download 2-Sheet Excel (.xlsx)",
            data=excel_bytes,
            file_name=export_filename,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )

# --- REINDEX LATERAL MATRIX (FULL DATES X SELECTED SELLERS) ---
if not df_display.empty or selected_sellers:
    df_display["day"] = pd.to_datetime(df_display["day"]).dt.strftime("%Y-%m-%d")

    all_days = sorted(pd.to_datetime(df_metric["day"]).dt.strftime("%Y-%m-%d").unique())

    if all_days and selected_sellers:
        full_grid = pd.MultiIndex.from_product(
            [selected_sellers, all_days],
            names=["company_seller_used_id", "day"]
        ).to_frame().reset_index(drop=True)

        df_display = pd.merge(
            full_grid,
            df_display,
            on=["company_seller_used_id", "day"],
            how="left"
        )

        df_display["qc_status"] = df_display["qc_status"].fillna("NO_DATA")
        df_display["collected_val"] = df_display["collected_val"].fillna(0)
        df_display["raw_val"] = df_display["raw_val"].fillna(0)
        df_display["accuracy_pct"] = df_display["accuracy_pct"].fillna(0)
        df_display["company_name"] = df_display["company_name"].fillna(selected_company)
        df_display["metric_name"] = df_display["metric_name"].fillna(selected_metric)

        if "seller_name" in df_metric.columns:
            id_to_name = dict(zip(df_metric["company_seller_used_id"], df_metric["seller_name"]))
            df_display["seller_name"] = df_display["company_seller_used_id"].map(id_to_name).fillna("")
        else:
            df_display["seller_name"] = ""

with st.container():
    st.markdown('<div class="sticky-header-anchor"></div>', unsafe_allow_html=True)

    total_sellers_count = df_company["company_seller_used_id"].nunique()

    st.markdown(
        f"""
        <div style="margin-bottom: 6px;">
            <div style="display: flex; justify-content: space-between; align-items: center;">
                <h3 style="margin: 0; font-size: 1.4rem; font-weight: 700; color: #f0f6fc;">🔥 QC Accuracy Dashboard</h3>
                <span style="font-size: 0.8rem; color: #8b949e;">
                    Snapshot: <code style="color: #3fb950; background: #161b22; padding: 2px 6px; border-radius: 4px;">{snapshot_time}</code>
                </span>
            </div>
            <div style="margin-top: 4px; font-size: 1.05rem; display: flex; gap: 12px; align-items: center;">
                <span style="color: #8b949e;">Company: 
                    <strong style="color: #58a6ff; background: #1f242c; padding: 3px 8px; border-radius: 6px; border: 1px solid #30363d;">
                        {selected_company}
                    </strong>
                </span>
                <span style="color: #8b949e;">Total Sellers: 
                    <strong style="color: #d2a8ff; background: #1f242c; padding: 3px 8px; border-radius: 6px; border: 1px solid #30363d;">
                        {total_sellers_count}
                    </strong>
                </span>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    total_count = len(df_display)
    good_count = (df_display["qc_status"] == "GOOD").sum()
    critical_count = (df_display["qc_status"] == "CRITICAL").sum()
    has_data_count = (df_display["qc_status"] == "HAS_DATA").sum()
    no_data_count = (df_display["qc_status"] == "NO_DATA").sum()

    good_pct = (good_count / total_count * 100) if total_count else 0
    critical_pct = (critical_count / total_count * 100) if total_count else 0
    has_data_pct = (has_data_count / total_count * 100) if total_count else 0
    no_data_pct = (no_data_count / total_count * 100) if total_count else 0

    col1, col2, col3, col4 = st.columns(4)
    col1.metric("🟩 GOOD (95-105%)", f"{good_count:,}", f"{good_pct:.1f}%", delta_color="off")
    col2.metric("🟥 CRITICAL", f"{critical_count:,}", f"{critical_pct:.1f}%", delta_color="off")
    col3.metric("🟦 HAS DATA", f"{has_data_count:,}", f"{has_data_pct:.1f}%", delta_color="off")
    col4.metric("⬛ NO DATA", f"{no_data_count:,}", f"{no_data_pct:.1f}%", delta_color="off")

    st.markdown(
        f"<h4 style='margin: 6px 0 0 0; font-size: 1.05rem; color: #c9d1d9;'>🔥 Heatmap Matrix: {selected_metric}</h4>",
        unsafe_allow_html=True)

if df_display.empty:
    st.info("No data matches the selected Seller filter.")
else:
    color_map = {
        "GOOD": "#2ea043",
        "CRITICAL": "#da3633",
        "HAS_DATA": "#1f6beb",
        "NO_DATA": "#21262d",
    }


    # Hiển thị rõ ràng seller_id kèm theo tên shop bên cạnh ngăn cách bằng khoảng trắng/dấu gạch nhẹ trên trục Y
    def make_chart_y_label(row):
        sid = row["company_seller_used_id"]
        sname = row.get("seller_name", "")
        if sname and pd.notna(sname) and str(sname).strip() != "":
            return f"{sid}   |   {sname}"
        return sid


    df_display["chart_y_label"] = df_display.apply(make_chart_y_label, axis=1)

    df_display = df_display.sort_values(by=["day", "company_seller_used_id"])

    fig = px.scatter(
        df_display,
        x="day",
        y="chart_y_label",
        color="qc_status",
        color_discrete_map=color_map,
        custom_data=[
            "company_seller_used_id",
            "seller_name",
            "collected_val",
            "raw_val",
            "accuracy_pct",
            "qc_status",
        ],
    )

    fig.update_traces(
        marker=dict(symbol="square", size=14),
        hovertemplate=(
            "<b>Seller ID</b>=%{customdata[0]}<br>"
            "<b>Seller Name</b>=%{customdata[1]}<br>"
            "<b>Date</b>=%{x}<br>"
            "<b>Database</b>=%{customdata[2]:,.0f}<br>"
            "<b>Raw Input</b>=%{customdata[3]:,.0f}<br>"
            "<b>Accuracy (%)</b>=%{customdata[4]:.2f}<br>"
            "<b>QC Status</b>=%{customdata[5]}"
            "<extra></extra>"
        ),
    )

    unique_sellers_raw = sorted(df_display["company_seller_used_id"].unique())
    id_to_chart_label = dict(zip(df_display["company_seller_used_id"], df_display["chart_y_label"]))
    unique_chart_labels = [id_to_chart_label[s] for s in unique_sellers_raw if s in id_to_chart_label]
    num_sellers = len(unique_sellers_raw)

    chart_height = max(350, num_sellers * 32)

    fig.update_layout(
        plot_bgcolor="#0d1117",
        paper_bgcolor="#0d1117",
        font=dict(color="#c9d1d9"),
        height=chart_height,
        showlegend=False,
        margin=dict(l=0, r=10, t=10, b=10),
        xaxis=dict(
            type="category",
            title="",
            tickangle=-45,
            gridcolor="#21262d",
            categoryorder="category descending",
        ),
        yaxis=dict(
            type="category",
            title="",
            gridcolor="#21262d",
            categoryorder="array",
            categoryarray=unique_chart_labels,
            autorange="reversed",
        ),
    )

    st.plotly_chart(fig, use_container_width=True, config={"displayModeBar": False})

with st.expander("📋 View Detailed Data Table (Critical Issues Only)"):
    df_critical = df_display[df_display["qc_status"] == "CRITICAL"].sort_values(
        by=["day", "company_seller_used_id"],
        ascending=[False, True]
    )

    if df_critical.empty:
        st.success("🎉 No CRITICAL issues found for the selected filter!")
    else:
        st.dataframe(
            df_critical[[
                "company_name", "company_seller_used_id", "seller_name", "day", "metric_name",
                "collected_val", "raw_val", "accuracy_pct", "qc_status"
            ]],
            use_container_width=True,
        )